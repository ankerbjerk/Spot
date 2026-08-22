#!/usr/bin/env python3
"""KiDiBOM — turn a KiCad BOM CSV into a DigiKey-priced XLSX.

Canonical source, updates and issues: https://github.com/ankerbjerk/KiDiBOM

This file is self-contained and safe to copy into other project repos as-is.
Run it with `python kidibom.py path/to/bom.csv` — no manual `pip install` is
required, it bootstraps its own dependencies into a local `.venv` on first run.

A local `.env` file is required in the same folder as the script, containing
the DigiKey API credentials — see the README for instructions on how to set it
up.
"""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

REQUIRED_PACKAGES = [
    "openpyxl>=3.1,<4",
    "requests>=2.31,<3",
    "python-dotenv>=1.0,<2",
]


def _venv_python(venv_dir: Path) -> Path:
    if os.name == "nt":
        return venv_dir / "Scripts" / "python.exe"
    return venv_dir / "bin" / "python"


def _ensure_deps_and_relaunch() -> None:
    """Make sure openpyxl/requests/python-dotenv are importable, bootstrapping
    a local venv and re-launching under it if they aren't. Must run before any
    third-party imports below."""
    try:
        import dotenv  # noqa: F401
        import openpyxl  # noqa: F401
        import requests  # noqa: F401

        return
    except ImportError:
        pass

    if os.environ.get("KIDIBOM_BOOTSTRAPPED") == "1":
        sys.stderr.write(
            "KiDiBOM: required packages are still missing inside the bootstrap "
            "virtual environment.\nInstall them manually and try again:\n"
            f"  pip install {' '.join(REQUIRED_PACKAGES)}\n"
        )
        sys.exit(1)

    script_path = Path(__file__).resolve()
    venv_dir = script_path.parent / ".venv"
    venv_python = _venv_python(venv_dir)

    try:
        if not venv_python.exists():
            print(
                "KiDiBOM: first run detected -- setting up a local virtual "
                "environment in .venv (this only happens once)...",
                file=sys.stderr,
            )
            subprocess.check_call([sys.executable, "-m", "venv", str(venv_dir)])
            subprocess.check_call(
                [str(venv_python), "-m", "pip", "install", *REQUIRED_PACKAGES]
            )
    except (subprocess.CalledProcessError, OSError) as exc:
        sys.stderr.write(
            f"KiDiBOM: automatic dependency setup failed ({exc}).\n"
            "Install the dependencies manually and re-run, e.g.:\n"
            f"  pip install {' '.join(REQUIRED_PACKAGES)}\n"
        )
        sys.exit(1)

    env = os.environ.copy()
    env["KIDIBOM_BOOTSTRAPPED"] = "1"
    try:
        os.execve(str(venv_python), [str(venv_python), str(script_path), *sys.argv[1:]], env)
    except OSError as exc:
        sys.stderr.write(f"KiDiBOM: failed to relaunch under the virtual environment ({exc}).\n")
        sys.exit(1)


_ensure_deps_and_relaunch()

# --------------------------------------------------------------------------
# Everything below can safely assume openpyxl / requests / python-dotenv are
# installed and importable.
# --------------------------------------------------------------------------

import argparse
import csv
import json
import logging
import re
import time
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import quote

import requests
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------

DIGIKEY_API_BASE = "https://api.digikey.com"
DIGIKEY_SANDBOX_API_BASE = "https://sandbox-api.digikey.com"
TOKEN_PATH = "/v1/oauth2/token"
PRODUCT_DETAILS_PATH = "/products/v4/search/{pn}/productdetails"
KEYWORD_SEARCH_PATH = "/products/v4/search/keyword"


@dataclass
class LocaleProfile:
    site: str
    language: str
    currency: str
    domain: str
    number_format: str


# NOK is the default (this script's primary user is in Norway); --usd switches
# to the US/dollar profile. DigiKey's API returns prices already denominated in
# the requested currency and a ProductUrl already pointing at the matching
# country site, so switching locale headers is enough -- no manual URL rewriting.
LOCALE_PROFILES: Dict[str, "LocaleProfile"] = {
    "nok": LocaleProfile(
        site="NO", language="EN", currency="NOK", domain="www.digikey.no", number_format='#,##0.00 "kr"'
    ),
    "usd": LocaleProfile(
        site="US", language="EN", currency="USD", domain="www.digikey.com", number_format='"$"#,##0.00'
    ),
}

ASSEMBLY_SHEET_NAME = "Assembly BOM"
GENERAL_SHEET_NAME = "General BOM"

BOARDS_CELL = "B2"
TOTAL_CELL = "B3"
HEADER_ROW = 5

NEW_COLUMNS = ["Order Qty", "Unit Cost", "Stock", "Status", "Qty x Cost", "Purchase URL", "Notes"]

REFERENCE_ALIASES = {"reference", "designator", "references", "designators"}
QTY_ALIASES = {"qty", "quantity"}
DIGIKEY_PN_ALIASES = {"digikey pn", "digi-key pn", "digikey part number", "digi-key part number"}
MPN_ALIASES = {"mpn", "manufacturer part number", "manufacturer pn"}

DEBUG_DUMP_PATH = Path("kidibom_digikey_debug.json")

# --------------------------------------------------------------------------
# Styling (Calibri ships with Excel and is substituted by an identical
# metric-compatible "Carlito" on LibreOffice/Linux, so layout is consistent
# cross-platform without requiring any specific font to be installed).
# --------------------------------------------------------------------------

FONT_NAME = "Calibri"
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True, size=11)
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
BODY_FONT = Font(name=FONT_NAME, size=11)
TOTAL_FONT = Font(name=FONT_NAME, bold=True, size=11)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
INFO_FILL = PatternFill("solid", fgColor="F2F2F2")

_THIN_SIDE = Side(style="thin", color="B7B7B7")
CELL_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

RIGHT_ALIGN = Alignment(horizontal="right")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")


def _style_row(ws: Worksheet, row: int, first_col: int, last_col: int, font: Font, fill: Optional[PatternFill] = None) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = font
        cell.border = CELL_BORDER
        if fill is not None:
            cell.fill = fill


def _autofit_column_widths(
    ws: Worksheet,
    headers: List[str],
    first_data_row: int,
    last_data_row: int,
    min_width: int = 10,
    max_width: int = 45,
) -> None:
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for r in range(first_data_row, last_data_row + 1):
            value = ws.cell(row=r, column=col_idx).value
            if value is None:
                continue
            text = str(value)
            if text.startswith("="):
                continue  # formula source text isn't representative of the displayed value
            max_len = max(max_len, len(text))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, min_width), max_width)


class ConfigError(Exception):
    """Raised for problems the user needs to fix (bad CSV, missing creds)."""


class DigiKeyLookupError(Exception):
    """Raised for a single failed DigiKey lookup; callers catch this per-row."""


# --------------------------------------------------------------------------
# CSV handling
# --------------------------------------------------------------------------


def _normalize_header_key(header: str) -> str:
    return re.sub(r"\s+", " ", header.strip().lower()).rstrip(".")


def normalize_headers(fieldnames: List[str]) -> Dict[str, str]:
    """Map canonical roles (reference/qty/digikey_pn/mpn) to the CSV's own
    header text, so we can look up values without forcing a rename."""
    role_map: Dict[str, str] = {}
    for name in fieldnames:
        key = _normalize_header_key(name)
        if key in REFERENCE_ALIASES:
            role_map.setdefault("reference", name)
        elif key in QTY_ALIASES:
            role_map.setdefault("qty", name)
        elif key in DIGIKEY_PN_ALIASES:
            role_map.setdefault("digikey_pn", name)
        elif key in MPN_ALIASES:
            role_map.setdefault("mpn", name)

    missing = [role for role in ("reference", "qty") if role not in role_map]
    if missing:
        raise ConfigError(
            f"Could not find required column(s) for {missing} in CSV headers {fieldnames}. "
            f"Expected a 'Reference'/'Designator' column and a 'Qty.'/'Qty'/'Quantity' column."
        )
    return role_map


def _coerce_qty(raw: str, reference: str) -> int:
    raw = (raw or "").strip()
    match = re.search(r"\d+", raw)
    if not match:
        logging.warning("Reference %s: could not parse quantity from %r; defaulting to 0", reference, raw)
        return 0
    return int(match.group())


@dataclass
class BomRow:
    original: Dict[str, str]
    reference: str
    qty: int
    digikey_pn: str
    mpn: str


def read_bom_rows(csv_path: Path) -> Tuple[List[str], Dict[str, str], List[BomRow]]:
    with csv_path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        role_map = normalize_headers(fieldnames)

        rows: List[BomRow] = []
        for raw in reader:
            reference = (raw.get(role_map["reference"]) or "").strip()
            qty = _coerce_qty(raw.get(role_map["qty"]) or "", reference)
            digikey_pn = (raw.get(role_map.get("digikey_pn", ""), "") or "").strip()
            mpn = (raw.get(role_map.get("mpn", ""), "") or "").strip()
            rows.append(
                BomRow(original=dict(raw), reference=reference, qty=qty, digikey_pn=digikey_pn, mpn=mpn)
            )
        return fieldnames, role_map, rows


# --------------------------------------------------------------------------
# DigiKey API client
# --------------------------------------------------------------------------


@dataclass
class PriceBreak:
    break_quantity: int
    unit_price: float


@dataclass
class DigiKeyResult:
    price_breaks: List[PriceBreak] = field(default_factory=list)
    stock: Optional[int] = None
    status: Optional[str] = None
    purchase_url: Optional[str] = None
    matched_pn: Optional[str] = None


def select_unit_price(breaks: List[PriceBreak], order_qty: int) -> Optional[float]:
    if not breaks:
        return None
    sorted_breaks = sorted(breaks, key=lambda b: b.break_quantity)
    candidates = [b for b in sorted_breaks if b.break_quantity <= order_qty]
    chosen = candidates[-1] if candidates else sorted_breaks[0]
    return chosen.unit_price


def _extract_status(product: Dict[str, Any]) -> Optional[str]:
    status = product.get("ProductStatus")
    if isinstance(status, dict):
        return status.get("Status")
    if isinstance(status, str):
        return status
    return None


def _extract_price_breaks(variation: Optional[Dict[str, Any]]) -> List[PriceBreak]:
    if not variation:
        return []
    breaks: List[PriceBreak] = []
    for entry in variation.get("StandardPricing") or []:
        try:
            breaks.append(PriceBreak(int(entry["BreakQuantity"]), float(entry["UnitPrice"])))
        except (KeyError, TypeError, ValueError):
            continue
    return breaks


def _select_variation(product: Dict[str, Any], requested_pn: Optional[str]) -> Optional[Dict[str, Any]]:
    variations = product.get("ProductVariations") or []
    if not variations:
        return None
    if requested_pn:
        for variation in variations:
            if str(variation.get("DigiKeyProductNumber", "")).strip().lower() == requested_pn.strip().lower():
                return variation
    return variations[0]


def _fallback_purchase_url(term: Optional[str], domain: str) -> Optional[str]:
    if not term:
        return None
    return f"https://{domain}/en/products/result?keywords={quote(term)}"


def parse_product_details(
    product: Dict[str, Any], domain: str, requested_pn: Optional[str] = None
) -> DigiKeyResult:
    variation = _select_variation(product, requested_pn)
    price_breaks = _extract_price_breaks(variation)

    stock = product.get("QuantityAvailable")
    if stock is None and variation:
        stock = variation.get("QuantityAvailableforPackageType")

    status = _extract_status(product)
    matched_pn = (variation or {}).get("DigiKeyProductNumber") or requested_pn

    # DigiKey's own ProductUrl is already localized to the requested site (it
    # respects the X-DIGIKEY-Locale-Site header), so it's used as-is when present.
    purchase_url = product.get("ProductUrl") or (variation or {}).get("ProductUrl")
    if not purchase_url:
        purchase_url = _fallback_purchase_url(matched_pn or requested_pn, domain)

    return DigiKeyResult(
        price_breaks=price_breaks,
        stock=stock,
        status=status,
        purchase_url=purchase_url,
        matched_pn=matched_pn,
    )


def select_search_candidate(
    products: List[Dict[str, Any]], requested_mpn: str
) -> Tuple[Optional[Dict[str, Any]], bool]:
    """Returns (product, inexact) — inexact is True if no exact MPN match was found."""
    for product in products:
        mpn_field = product.get("ManufacturerProductNumber") or product.get("ManufacturerPartNumber")
        if mpn_field and str(mpn_field).strip().lower() == requested_mpn.strip().lower():
            return product, False
    if products:
        return products[0], True
    return None, False


class DigiKeyClient:
    def __init__(
        self,
        client_id: str,
        client_secret: str,
        locale: LocaleProfile,
        sandbox: bool = False,
        debug: bool = False,
    ):
        self.client_id = client_id
        self.client_secret = client_secret
        self.locale = locale
        self.base_url = DIGIKEY_SANDBOX_API_BASE if sandbox else DIGIKEY_API_BASE
        self.debug = debug
        self._debug_dumped = False
        self.session = requests.Session()
        self._token: Optional[str] = None
        self._token_expiry = 0.0

    def _ensure_token(self) -> None:
        if self._token and time.monotonic() < self._token_expiry:
            return
        response = self.session.post(
            self.base_url + TOKEN_PATH,
            data={
                "client_id": self.client_id,
                "client_secret": self.client_secret,
                "grant_type": "client_credentials",
            },
            timeout=15,
        )
        response.raise_for_status()
        payload = response.json()
        self._token = payload["access_token"]
        self._token_expiry = time.monotonic() + float(payload.get("expires_in", 599)) - 30

    def _headers(self) -> Dict[str, str]:
        self._ensure_token()
        return {
            "Authorization": f"Bearer {self._token}",
            "X-DIGIKEY-Client-Id": self.client_id,
            "X-DIGIKEY-Locale-Site": self.locale.site,
            "X-DIGIKEY-Locale-Language": self.locale.language,
            "X-DIGIKEY-Locale-Currency": self.locale.currency,
            "Accept": "application/json",
        }

    def _maybe_dump_debug(self, data: Any) -> None:
        if not self.debug or self._debug_dumped:
            return
        self._debug_dumped = True
        try:
            DEBUG_DUMP_PATH.write_text(json.dumps(data, indent=2), encoding="utf-8")
            logging.info("Dumped first raw DigiKey response to %s", DEBUG_DUMP_PATH)
        except OSError as exc:
            logging.warning("Could not write DigiKey debug dump: %s", exc)

    def _request_with_retry(self, method: str, url: str, **kwargs: Any) -> requests.Response:
        last_error: Optional[BaseException] = None
        for attempt in range(2):
            try:
                response = self.session.request(method, url, headers=self._headers(), timeout=20, **kwargs)
            except requests.RequestException as exc:
                last_error = exc
                time.sleep(1.5 * (attempt + 1))
                continue
            if response.status_code == 429 or response.status_code >= 500:
                last_error = DigiKeyLookupError(f"HTTP {response.status_code} from DigiKey")
                time.sleep(1.5 * (attempt + 1))
                continue
            return response
        raise DigiKeyLookupError(str(last_error) if last_error else "DigiKey request failed")

    def get_product_details(self, digikey_pn: str) -> Dict[str, Any]:
        url = f"{self.base_url}{PRODUCT_DETAILS_PATH.format(pn=quote(digikey_pn, safe=''))}"
        response = self._request_with_retry("GET", url)
        if response.status_code == 404:
            raise DigiKeyLookupError(f"DigiKey PN not found: {digikey_pn}")
        response.raise_for_status()
        data = response.json()
        self._maybe_dump_debug(data)
        # The productdetails endpoint wraps the actual product under a "Product"
        # key rather than returning it at the response root.
        return data.get("Product", data)

    def search_by_keyword(self, mpn: str) -> Dict[str, Any]:
        url = f"{self.base_url}{KEYWORD_SEARCH_PATH}"
        response = self._request_with_retry("POST", url, json={"Keywords": mpn, "Limit": 10})
        response.raise_for_status()
        data = response.json()
        self._maybe_dump_debug(data)
        return data


def lookup_digikey_cached(
    client: DigiKeyClient,
    cache: Dict[Tuple[str, str], Optional[DigiKeyResult]],
    reference: str,
    pn: Optional[str] = None,
    mpn: Optional[str] = None,
) -> Optional[DigiKeyResult]:
    if pn:
        key = ("pn", pn.strip().lower())
    elif mpn:
        key = ("mpn", mpn.strip().lower())
    else:
        return None

    if key in cache:
        return cache[key]

    result: Optional[DigiKeyResult] = None
    try:
        if pn:
            data = client.get_product_details(pn)
            result = parse_product_details(data, client.locale.domain, requested_pn=pn)
        else:
            assert mpn is not None
            data = client.search_by_keyword(mpn)
            products = data.get("Products") or data.get("products") or []
            product, inexact = select_search_candidate(products, mpn)
            if product is None:
                raise DigiKeyLookupError(f"No DigiKey results found for MPN {mpn!r}")
            result = parse_product_details(product, client.locale.domain)
            if inexact:
                logging.info(
                    "Reference %s: no exact DigiKey match for MPN %r; using closest result %s",
                    reference,
                    mpn,
                    result.matched_pn,
                )
    except DigiKeyLookupError as exc:
        logging.warning("Reference %s: DigiKey lookup failed for %s (%s)", reference, pn or mpn, exc)
    except requests.RequestException as exc:
        logging.warning("Reference %s: DigiKey request error for %s (%s)", reference, pn or mpn, exc)

    cache[key] = result
    return result


# --------------------------------------------------------------------------
# Building output rows
# --------------------------------------------------------------------------


@dataclass
class OutputRow:
    original: Dict[str, str]
    qty: int
    order_qty: int
    unit_price: Optional[float] = None
    stock: Optional[int] = None
    status: Optional[str] = None
    line_total: Optional[float] = None
    purchase_url: Optional[str] = None
    note: Optional[str] = None


def build_output_rows(
    bom_rows: List[BomRow], boards: int, client: Optional[DigiKeyClient]
) -> Tuple[List[OutputRow], float]:
    cache: Dict[Tuple[str, str], Optional[DigiKeyResult]] = {}
    output_rows: List[OutputRow] = []
    total = 0.0

    for row in bom_rows:
        order_qty = row.qty * boards
        out_row = OutputRow(original=row.original, qty=row.qty, order_qty=order_qty)

        result = None
        if client is not None and (row.digikey_pn or row.mpn):
            result = lookup_digikey_cached(
                client, cache, row.reference, pn=row.digikey_pn or None, mpn=row.mpn or None
            )
            if result is None:
                out_row.note = "DigiKey lookup failed (see console warnings)"

        if result is not None:
            out_row.stock = result.stock
            out_row.status = result.status
            out_row.purchase_url = result.purchase_url
            out_row.unit_price = select_unit_price(result.price_breaks, order_qty)
            if out_row.unit_price is not None:
                out_row.line_total = out_row.unit_price * order_qty
                total += out_row.line_total
            elif result.stock == 0:
                out_row.note = "No DigiKey pricing available (0 in stock)"
            else:
                out_row.note = "No DigiKey pricing available"

        output_rows.append(out_row)

    return output_rows, total


# --------------------------------------------------------------------------
# XLSX output
# --------------------------------------------------------------------------


def write_assembly_sheet(
    wb: Workbook,
    original_fieldnames: List[str],
    qty_field: str,
    output_rows: List[OutputRow],
    boards: int,
    number_format: str,
) -> Worksheet:
    if ASSEMBLY_SHEET_NAME in wb.sheetnames:
        del wb[ASSEMBLY_SHEET_NAME]
    ws = wb.create_sheet(ASSEMBLY_SHEET_NAME, 0)

    ws["A1"] = "KiDiBOM - Assembly BOM"
    ws["A1"].font = TITLE_FONT

    for coord in ("A2", "B2", "A3", "B3"):
        ws[coord].fill = INFO_FILL
        ws[coord].border = CELL_BORDER
    ws["A2"] = "No. boards:"
    ws["A2"].font = LABEL_FONT
    ws[BOARDS_CELL] = boards
    ws[BOARDS_CELL].font = LABEL_FONT

    ws["A3"] = "Total assembly cost:"
    ws["A3"].font = LABEL_FONT

    headers = list(original_fieldnames) + NEW_COLUMNS
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = CELL_BORDER
    ws.row_dimensions[HEADER_ROW].height = 20

    n_original = len(original_fieldnames)
    qty_col = original_fieldnames.index(qty_field) + 1
    order_qty_col = n_original + NEW_COLUMNS.index("Order Qty") + 1
    unit_cost_col = n_original + NEW_COLUMNS.index("Unit Cost") + 1
    stock_col = n_original + NEW_COLUMNS.index("Stock") + 1
    status_col = n_original + NEW_COLUMNS.index("Status") + 1
    qty_cost_col = n_original + NEW_COLUMNS.index("Qty x Cost") + 1
    purchase_url_col = n_original + NEW_COLUMNS.index("Purchase URL") + 1
    notes_col = n_original + NEW_COLUMNS.index("Notes") + 1

    qty_letter = get_column_letter(qty_col)
    order_qty_letter = get_column_letter(order_qty_col)
    unit_cost_letter = get_column_letter(unit_cost_col)

    numeric_cols = {qty_col, order_qty_col, unit_cost_col, stock_col, qty_cost_col}

    datasheet_field = next((f for f in original_fieldnames if _normalize_header_key(f) == "datasheet"), None)
    datasheet_col = original_fieldnames.index(datasheet_field) + 1 if datasheet_field else None
    hyperlink_cols = {purchase_url_col} | ({datasheet_col} if datasheet_col else set())

    data_start = HEADER_ROW + 1
    row_idx = data_start
    for out_row in output_rows:
        for col_idx, field_name in enumerate(original_fieldnames, start=1):
            if field_name == qty_field:
                ws.cell(row=row_idx, column=col_idx, value=out_row.qty)
            else:
                value = out_row.original.get(field_name, "") or ""
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if col_idx == datasheet_col and isinstance(value, str) and value.startswith("https://"):
                    cell.hyperlink = value
                    cell.style = "Hyperlink"

        # Order Qty and Qty x Cost are live formulas driven by No. boards (fixed
        # cell B2), so editing boards afterwards rescales quantities/totals
        # instantly. Unit Cost stays a fixed value from the price break that was
        # selected when the script ran -- it does not re-select a cheaper/pricier
        # DigiKey quantity break on its own; re-run the script for that.
        ws.cell(row=row_idx, column=order_qty_col, value=f"={qty_letter}{row_idx}*$B$2")
        ws.cell(row=row_idx, column=unit_cost_col, value=out_row.unit_price)
        ws.cell(row=row_idx, column=stock_col, value=out_row.stock)
        ws.cell(row=row_idx, column=status_col, value=out_row.status)
        ws.cell(
            row=row_idx,
            column=qty_cost_col,
            value=(
                f'=IF({unit_cost_letter}{row_idx}="","",'
                f"{order_qty_letter}{row_idx}*{unit_cost_letter}{row_idx})"
            ),
        )

        url_cell = ws.cell(row=row_idx, column=purchase_url_col)
        if out_row.purchase_url:
            url_cell.value = "View on DigiKey"
            url_cell.hyperlink = out_row.purchase_url
            url_cell.style = "Hyperlink"

        ws.cell(row=row_idx, column=notes_col, value=out_row.note)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = CELL_BORDER
            if col_idx not in hyperlink_cols:
                cell.font = BODY_FONT
            if col_idx in numeric_cols:
                cell.alignment = RIGHT_ALIGN

        row_idx += 1

    total_row_idx = row_idx
    qty_cost_letter = get_column_letter(qty_cost_col)
    if output_rows:
        total_value: Any = f"=SUM({qty_cost_letter}{data_start}:{qty_cost_letter}{total_row_idx - 1})"
    else:
        total_value = 0
    ws.cell(row=total_row_idx, column=1, value="Total")
    ws.cell(row=total_row_idx, column=qty_cost_col, value=total_value)
    _style_row(ws, total_row_idx, 1, len(headers), TOTAL_FONT, TOTAL_FILL)
    ws.cell(row=total_row_idx, column=qty_cost_col).number_format = number_format
    ws.cell(row=total_row_idx, column=qty_cost_col).alignment = RIGHT_ALIGN

    ws[TOTAL_CELL] = f"={qty_cost_letter}{total_row_idx}"
    ws[TOTAL_CELL].number_format = number_format

    for r in range(data_start, total_row_idx + 1):
        ws.cell(row=r, column=unit_cost_col).number_format = number_format
        ws.cell(row=r, column=qty_cost_col).number_format = number_format

    ws.freeze_panes = ws.cell(row=data_start, column=1).coordinate

    _autofit_column_widths(ws, headers, HEADER_ROW, total_row_idx)

    return ws


def ensure_general_bom_sheet(wb: Workbook, number_format: str) -> None:
    if GENERAL_SHEET_NAME in wb.sheetnames:
        return

    ws = wb.create_sheet(GENERAL_SHEET_NAME)
    ws["A1"] = "General BOM"
    ws["A1"].font = TITLE_FONT

    header_row = 3
    headers = ["Item", "Qty", "Unit Cost", "Line Total", "Notes"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = CELL_BORDER
    ws.row_dimensions[header_row].height = 20

    seed_row = header_row + 1
    ws.cell(row=seed_row, column=1, value="PCB Assembly (see 'Assembly BOM' sheet)")
    ws.cell(row=seed_row, column=2, value=f"='{ASSEMBLY_SHEET_NAME}'!{BOARDS_CELL}")
    ws.cell(row=seed_row, column=4, value=f"='{ASSEMBLY_SHEET_NAME}'!{TOTAL_CELL}")
    ws.cell(row=seed_row, column=5, value="Auto-synced from Assembly BOM -- edit costs there, not here")

    blank_rows = 10
    last_row = seed_row + blank_rows
    totals_row = last_row + 1

    # Pre-fill Line Total = Qty x Unit Cost for the blank rows the user fills in
    # by hand. The seed row above is left alone -- its Line Total is a live
    # cross-sheet reference to the Assembly BOM total, not a Qty x Cost product.
    for r in range(seed_row + 1, last_row + 1):
        ws.cell(row=r, column=4, value=f"=B{r}*C{r}")

    numeric_cols = {2, 3, 4}
    for r in range(seed_row, last_row + 1):
        ws.cell(row=r, column=3).number_format = number_format
        ws.cell(row=r, column=4).number_format = number_format
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col_idx)
            cell.font = BODY_FONT
            cell.border = CELL_BORDER
            if col_idx in numeric_cols:
                cell.alignment = RIGHT_ALIGN

    ws.cell(row=totals_row, column=1, value="Grand Total")
    total_cell = ws.cell(row=totals_row, column=4, value=f"=SUBTOTAL(109,D{seed_row}:D{last_row})")
    total_cell.number_format = number_format
    _style_row(ws, totals_row, 1, len(headers), TOTAL_FONT, TOTAL_FILL)
    ws.cell(row=totals_row, column=4).alignment = RIGHT_ALIGN

    table_ref = f"A{header_row}:E{last_row}"
    table = Table(displayName="GeneralBOM", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    _autofit_column_widths(ws, headers, header_row, totals_row, min_width=14)


# --------------------------------------------------------------------------
# CLI / orchestration
# --------------------------------------------------------------------------


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Enrich a KiCad BOM CSV with DigiKey cost/stock/status data and produce an XLSX BOM."
    )
    parser.add_argument("csv_path", help="Path to the KiCad-generated BOM CSV file")
    parser.add_argument(
        "-n", "--boards", type=int, default=5, help="Number of boards to order for (default: 5)"
    )
    parser.add_argument("--sandbox", action="store_true", help="Use the DigiKey sandbox API instead of production")
    parser.add_argument(
        "--usd",
        action="store_true",
        help=(
            "Price in US dollars via the digikey.com/en site instead of the default "
            "Norwegian kroner (NOK) via digikey.no/en"
        ),
    )
    parser.add_argument(
        "--debug-digikey", action="store_true", help="Dump the first raw DigiKey API response to a debug file"
    )
    parser.add_argument(
        "--env-file",
        default=None,
        help="Path to a .env file with DIGIKEY_CLIENT_ID/DIGIKEY_CLIENT_SECRET (default: .env next to this script)",
    )
    return parser.parse_args(argv)


def load_digikey_credentials(env_file: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    if env_file:
        load_dotenv(env_file)
    else:
        load_dotenv(Path(__file__).resolve().parent / ".env")
    return os.environ.get("DIGIKEY_CLIENT_ID"), os.environ.get("DIGIKEY_CLIENT_SECRET")


def main(argv: Optional[List[str]] = None) -> None:
    args = parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    csv_path = Path(args.csv_path).resolve()
    if not csv_path.is_file():
        logging.error("Input CSV not found: %s", csv_path)
        sys.exit(1)

    try:
        fieldnames, role_map, bom_rows = read_bom_rows(csv_path)
    except ConfigError as exc:
        logging.error(str(exc))
        sys.exit(1)

    needs_digikey = any(row.digikey_pn or row.mpn for row in bom_rows)
    locale = LOCALE_PROFILES["usd" if args.usd else "nok"]

    client: Optional[DigiKeyClient] = None
    if needs_digikey:
        client_id, client_secret = load_digikey_credentials(args.env_file)
        if not client_id or not client_secret:
            logging.error(
                "DigiKey API credentials not found. Set DIGIKEY_CLIENT_ID and DIGIKEY_CLIENT_SECRET "
                "in a .env file (see .env.example) or as environment variables. Register a free app "
                "at https://developer.digikey.com to obtain them."
            )
            sys.exit(1)
        client = DigiKeyClient(client_id, client_secret, locale, sandbox=args.sandbox, debug=args.debug_digikey)

    output_rows, total = build_output_rows(bom_rows, args.boards, client)

    out_path = csv_path.with_suffix(".xlsx")
    if out_path.exists():
        wb = load_workbook(out_path)
        if (
            "Sheet" in wb.sheetnames
            and len(wb.sheetnames) == 1
            and wb["Sheet"].max_row == 1
            and wb["Sheet"].max_column == 1
            and wb["Sheet"]["A1"].value is None
        ):
            del wb["Sheet"]
    else:
        wb = Workbook()
        wb.remove(wb.active)

    write_assembly_sheet(wb, fieldnames, role_map["qty"], output_rows, args.boards, locale.number_format)
    ensure_general_bom_sheet(wb, locale.number_format)

    wb.save(out_path)
    logging.info("Wrote %s (total assembly cost: %.2f)", out_path, total)


if __name__ == "__main__":
    main()
